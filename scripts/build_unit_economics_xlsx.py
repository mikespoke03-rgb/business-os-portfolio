#!/usr/bin/env python3
"""Build a channel unit-economics workbook with live Excel formulas (openpyxl).

Same structure as the research models built for the real project (one Inputs sheet with named, yellow assumption
cells; one sheet per channel; a comparison sheet), but every number here is ILLUSTRATIVE and synthetic.
Tax and customs treatment is deliberately simplified - verify with a professional before using it for decisions.

Usage: python3 scripts/build_unit_economics_xlsx.py [--out examples/unit_economics_example.xlsx]
"""
import argparse, datetime as dt, io, os, re, zipfile
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
B = Font(bold=True)
Y = PatternFill("solid", fgColor="FFF2CC")   # assumption - change me
G = PatternFill("solid", fgColor="E2EFDA")   # result

INPUTS = [  # key, value, description, is_assumption
    ("supplier_usd", 4.80, "Supplier price per unit, USD (synthetic)", 1),
    ("usd_rub", 90.0, "USD/RUB rate used in the model (illustrative)", 1),
    ("usd_local", 480.0, "USD to foreign-entity local currency rate (illustrative)", 1),
    ("fx_fees", 0.03, "Bank + payment agent + FX markup, share of purchase", 1),
    ("freight_ru_rub", 45.0, "Freight to RU warehouse per unit, RUB", 1),
    ("duty", 0.065, "Import duty rate (illustrative)", 1),
    ("vat_ru", 0.22, "VAT rate RU (verify current law)", 0),
    ("pack_rub", 25.0, "Packaging / labelling per unit, RUB", 1),
    ("ru_price", 1990.0, "RU marketplace seller price, RUB", 1),
    ("ru_comm", 0.28, "RU marketplace commission (FBO), share", 1),
    ("ru_logi", 95.0, "RU marketplace logistics per unit, RUB", 1),
    ("acq", 0.015, "Acquiring / payment processing, share", 1),
    ("ads", 0.08, "Ads cost, share of price (CPC + pay-per-order)", 1),
    ("vat_base_share", 0.50, "Share of seller price used as VAT base in the RU model", 1),
    ("fe_price_rub", 1990.0, "Foreign entity -> RU buyer: seller price, RUB", 1),
    ("fe_comm", 0.12, "Marketplace commission under the foreign-seller contract, share", 1),
    ("fe_logi_rub", 140.0, "Logistics foreign warehouse -> RU buyer per unit, RUB", 1),
    ("fe_freight_rub", 38.0, "Freight to foreign warehouse per unit, RUB", 1),
    ("fe_vat_import", 0.16, "Foreign import VAT (assumed creditable -> cash-flow cost only)", 1),
    ("fe_ru_vat_on_top", 0.22, "RU VAT charged on top of the price for a foreign seller (verify)", 1),
    ("fe_fixed_per_unit_rub", 22.0, "Certification + broker + local entity costs, amortised per unit, RUB", 1),
    ("cb_price_rub", 1990.0, "Cross-border seller -> RU buyer: price, RUB", 1),
    ("cb_platform_fee", 0.18, "Cross-border platform fee, share", 1),
    ("cb_parcel_rub", 260.0, "International parcel per unit, RUB", 1),
    ("cb_duty_share", 0.0, "Share of orders above the duty-free threshold", 1),
    ("own_price_rub", 1790.0, "Own web store price, RUB", 1),
    ("own_delivery_rub", 160.0, "Delivery to pick-up point per order, RUB", 1),
    ("own_return_share", 0.05, "Non-buyout / return share", 1),
    ("own_cac_1", 0.0, "CAC scenario 1, RUB per order", 1),
    ("own_cac_2", 300.0, "CAC scenario 2, RUB per order", 1),
    ("own_cac_3", 600.0, "CAC scenario 3, RUB per order", 1),
    ("target_margin", 0.15, "Target margin share used for the 'price needed' line", 1),
]


def sheet(wb, title, rows, note):
    ws = wb.create_sheet(title)
    ws.append([note]); ws["A1"].font = B
    ws.append(["Line", "RUB per unit", "Formula logic"])
    for c in ws[2]:
        c.font = B
    for i, (label, formula, logic) in enumerate(rows, start=3):
        ws.cell(i, 1, label); ws.cell(i, 2, formula); ws.cell(i, 3, logic)
        if label.upper().startswith("MARGIN") or label.startswith("Price needed"):
            ws.cell(i, 1).font = B; ws.cell(i, 2).fill = G
    ws.column_dimensions["A"].width = 52; ws.column_dimensions["B"].width = 16; ws.column_dimensions["C"].width = 70
    return ws


FIXED_TIME = dt.datetime(2026, 1, 1, 0, 0, 0)   # reproducible file: no build time in metadata or zip entries


def save_deterministic(wb, out):
    """openpyxl stamps the current time into docProps and zip entries; pin both so CI can diff examples/."""
    wb.properties.creator = "DemoBrand synthetic example"
    wb.properties.created = FIXED_TIME
    wb.properties.modified = FIXED_TIME
    buf = io.BytesIO(); wb.save(buf)
    src = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == "docProps/core.xml":
                text = data.decode("utf-8")
                for tag in ("created", "modified"):
                    text = re.sub(rf"(<dcterms:{tag}[^>]*>)[^<]*(</dcterms:{tag}>)", rf"\g<1>2026-01-01T00:00:00Z\g<2>", text)
                data = text.encode("utf-8")
            fixed = zipfile.ZipInfo(info.filename, date_time=(2026, 1, 1, 0, 0, 0))
            fixed.compress_type = zipfile.ZIP_DEFLATED
            fixed.external_attr = 0o644 << 16
            dst.writestr(fixed, data)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default=os.path.join(ROOT, "examples", "unit_economics_example.xlsx"))
    a = ap.parse_args()
    wb = Workbook()
    rd = wb.active; rd.title = "README"
    for line in ["DemoBrand - unit economics by sales channel (ILLUSTRATIVE, synthetic numbers).",
                 "Yellow cells on 'Inputs' are assumptions; everything else recalculates with formulas.",
                 "Margin = per bought-out unit, before profit tax. Tax/customs logic is simplified - verify before use.",
                 "Sheets: RU_FBO (base), Foreign_entity (a foreign company sells to RU buyers), Crossborder, OwnStore (3 CAC scenarios), Compare.",
                 "The Compare sheet only reflects the synthetic inputs - it is not a conclusion about any real channel."]:
        rd.append([line])
    rd.column_dimensions["A"].width = 120
    inp = wb.create_sheet("Inputs")
    inp.append(["key", "value", "description"])
    for c in inp[1]:
        c.font = B
    for r, (k, v, d, ass) in enumerate(INPUTS, start=2):
        inp.cell(r, 1, k); c = inp.cell(r, 2, v); inp.cell(r, 3, d)
        if ass:
            c.fill = Y
        wb.defined_names[k] = DefinedName(k, attr_text=f"Inputs!$B${r}")
    inp.column_dimensions["A"].width = 22; inp.column_dimensions["C"].width = 80

    # landed cost block reused by every channel
    landed = [
        ("Purchase, RUB", "=supplier_usd*usd_rub*(1+fx_fees)", "supplier price x FX x (1 + bank/agent/FX fees)"),
        ("Freight to RU", "=freight_ru_rub", "input"),
        ("Duty", "=(B3+B4)*duty", "(purchase + freight) x duty"),
        ("Packaging", "=pack_rub", "input"),
        ("COGS ex import VAT", "=B3+B4+B5+B6", "sum"),
        ("Import VAT (creditable)", "=(B3+B4+B5)*vat_ru", "customs value x VAT; reduces VAT payable"),
    ]
    sheet(wb, "RU_FBO", landed + [
        ("Seller price", "=ru_price", "input"),
        ("Marketplace commission", "=B9*ru_comm", "price x commission"),
        ("Logistics", "=ru_logi", "input"),
        ("Acquiring", "=B9*acq", "price x rate"),
        ("Ads", "=B9*ads", "price x ads share"),
        ("VAT payable (model)", "=B9*vat_base_share*vat_ru/(1+vat_ru)-B8", "price x base share x 22/122 - import VAT"),
        ("MARGIN, RUB", "=B9-B10-B11-B12-B13-B14-B7", "price - all lines - COGS"),
        ("MARGIN, % of price", "=B15/B9", ""),
        ("Price needed for target margin", "=(B11+B7-B8)/(1-ru_comm-acq-ads-vat_base_share*vat_ru/(1+vat_ru)-target_margin)", "solves margin = target x price"),
    ], "RU marketplace, FBO (base case)")
    sheet(wb, "Foreign_entity", [
        ("Purchase, RUB", "=supplier_usd*usd_rub*(1+fx_fees)", "same supplier price"),
        ("Freight to foreign warehouse", "=fe_freight_rub", "input"),
        ("Duty (foreign)", "=(B3+B4)*duty", "illustrative common tariff"),
        ("Packaging", "=pack_rub", "input"),
        ("COGS", "=B3+B4+B5+B6", "sum"),
        ("Foreign import VAT (cash-flow only)", "=(B3+B4+B5)*fe_vat_import", "assumed creditable; shown, not deducted"),
        ("Seller price", "=fe_price_rub", "input"),
        ("RU VAT on top (paid by buyer, remitted)", "=B9*fe_ru_vat_on_top", "not seller income"),
        ("Marketplace commission", "=B9*fe_comm", "price x commission"),
        ("Logistics foreign warehouse -> RU buyer", "=fe_logi_rub", "input"),
        ("Acquiring + ads", "=B9*(acq+ads)", ""),
        ("Entity / certification costs per unit", "=fe_fixed_per_unit_rub", "amortised"),
        ("MARGIN, RUB", "=B9-B11-B12-B13-B14-B7", "price - lines - COGS"),
        ("MARGIN, % of price", "=B15/B9", ""),
    ], "Foreign entity selling to RU buyers through the marketplace foreign-seller contract (simplified)")
    sheet(wb, "Crossborder", [
        ("Purchase, RUB", "=supplier_usd*usd_rub*(1+fx_fees)", "same supplier price"),
        ("International parcel", "=cb_parcel_rub", "per unit"),
        ("Expected duty", "=cb_duty_share*B9*duty", "only orders above the threshold"),
        ("Packaging", "=pack_rub", "input"),
        ("COGS (landed at buyer)", "=B3+B4+B6", "duty kept separate"),
        ("(no RU import VAT for the seller)", 0, "cross-border parcel model"),
        ("Price", "=cb_price_rub", "input"),
        ("Platform fee", "=B9*cb_platform_fee", "price x fee"),
        ("Acquiring + ads", "=B9*(acq+ads)", ""),
        ("MARGIN, RUB", "=B9-B10-B11-B5-B7", "price - fee - acq/ads - duty - COGS"),
        ("MARGIN, % of price", "=B12/B9", ""),
    ], "Foreign company, cross-border parcels to RU buyers (simplified)")
    own = landed + [
        ("Price (buyer pays)", "=own_price_rub", "input"),
        ("Acquiring", "=B9*acq", ""),
        ("Delivery to pick-up point", "=own_delivery_rub", "seller pays delivery"),
        ("Returns: lost delivery on non-buyouts", "=own_return_share/(1-own_return_share)*own_delivery_rub", "per bought-out unit"),
        ("VAT payable", "=B9*vat_ru/(1+vat_ru)-B8", "full price is the VAT base"),
        ("Margin before CAC", "=B9-B10-B11-B12-B13-B7", ""),
        ("MARGIN, CAC scenario 1", "=B14-own_cac_1", ""),
        ("MARGIN, CAC scenario 2", "=B14-own_cac_2", ""),
        ("MARGIN, CAC scenario 3", "=B14-own_cac_3", ""),
    ]
    sheet(wb, "OwnStore", own, "Own web store + marketplace delivery service (stock stays at the FBO warehouse)")
    cmp = wb.create_sheet("Compare")
    cmp.append(["Channel", "Price, RUB", "Margin, RUB", "Margin, %"])
    for c in cmp[1]:
        c.font = B
    for ch, p, m in [("RU FBO", "RU_FBO!B9", "RU_FBO!B15"), ("Foreign entity", "Foreign_entity!B9", "Foreign_entity!B15"),
                     ("Cross-border", "Crossborder!B9", "Crossborder!B12"), ("Own store, CAC 2", "OwnStore!B9", "OwnStore!B16")]:
        r = cmp.max_row + 1
        cmp.cell(r, 1, ch); cmp.cell(r, 2, f"={p}"); cmp.cell(r, 3, f"={m}"); cmp.cell(r, 4, f"=C{r}/B{r}").fill = G
    cmp.column_dimensions["A"].width = 22
    save_deterministic(wb, a.out)
    print("saved", a.out)


if __name__ == "__main__":
    main()
